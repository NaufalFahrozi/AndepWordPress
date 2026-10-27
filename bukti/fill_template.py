import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

SRC = r'C:\Users\naufa\Downloads\Template Laporan UTS Praktik WordPress Andep.docx'
OUT = r'C:\Kuliah\AndepWordPress\bukti\UTS_CMS_2511102414081_Muhammad_Andep.docx'
IMG = r'C:\Kuliah\AndepWordPress\bukti\report'

doc = docx.Document(SRC)

def set_para_text(p, text):
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in p.runs[1:]:
        r.text = ''

paras = doc.paragraphs
# Fill identity fields, preserving template run formatting
set_para_text(paras[2], 'Nama: Muhammad Andep          NIM: 2511102414081')
set_para_text(paras[3], 'Kelas: -          Varian soal (A/B/C/D): C')
set_para_text(paras[4], 'Nama website: Studio Karya Nusa')
set_para_text(paras[12], 'Nama: Muhammad Andep          NIM: 2511102414081')

# Final checklist: mark all as checked
set_para_text(paras[17], 'Pemeriksaan akhir:  [v] 2 Page terbit    [v] 2 Post terbit    [v] Category & Tag benar    [v] Gambar tampil    [v] Pengaturan tersimpan    [v] Kode NIM sesuai')
set_para_text(paras[18], 'Nama berkas pengumpulan: UTS_CMS_2511102414081_Muhammad_Andep.pdf. Hasil di WordPress tetap disimpan untuk pemeriksaan dosen.')

# Replace the six placeholder cells with screenshots + captions
captions = [
    'Page 1: Profil Studio Karya Nusa  |  Page 2: Program Studio Karya Nusa',
    'Post 1: Galeri Poster Komunitas  |  Post 2: Lokakarya Poster Digital',
    'Karya Komunitas / Poster  |  Agenda Kreatif / Desain Digital',
    'Media Library: c_studio.jpg  |  Gambar disisipkan di Post 1',
    'Site Title: Studio Karya Nusa  |  Tagline: Ruang kreasi dan berbagi ide',
    'Lokasi: Samarinda Ulu  |  Jadwal: Jumat, 14.00 WITA  |  Kode: UTS-C-2511102414081'
]

for i, table in enumerate(doc.tables):
    cell = table.rows[1].cells[0]
    # Clear cell content while keeping one paragraph
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Remove spacing that could waste vertical room
    p.paragraph_format.space_before = 0
    p.paragraph_format.space_after = 0
    run = p.add_run()
    img_path = os.path.join(IMG, f'B{i+1}.png')
    run.add_picture(img_path, width=Inches(6.0))
    cap = cell.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(1)
    cap.paragraph_format.space_after = Pt(1)
    cr = cap.add_run(captions[i])
    cr.font.name = 'Arial'
    cr.font.size = Pt(7.5)
    cr.font.color.rgb = RGBColor(0x5E, 0x6A, 0x63)
    print('inserted', img_path)

doc.save(OUT)
print('saved', OUT)

